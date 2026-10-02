from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.py_correction.engine.genote.internal.dtos import EvaluationInfo, StudentInfo
from src.py_correction.engine.helpers import decimal_helpers


class GeNoteGradeSheetParser:

    def __init__(self, file_path: str | Path):
        self.file_path = file_path
        self.wb = load_workbook(self.file_path, data_only=True)
        self.ws = self.wb[self.wb.sheetnames[0]]

    @property
    def get_course_name(self) -> str:
        course_name = self._find_value_by_text_offset("Cours:", col_offset=1)
        return course_name.strip() if course_name is not None else ""

    @property
    def get_student_list(self) -> list[StudentInfo]:
        cip_list = self._get_value_list("Cip", row_offset=1)
        student_name_list = self._get_value_list("Cip", row_offset=1, col_offset=1)
        student_cip_dict = dict(zip(cip_list, student_name_list, strict=True))

        return [StudentInfo(cip=cip,
                            first_name=student_name.split(",")[1].strip(),
                            last_name=student_name.split(",")[0].strip())
                for cip, student_name in student_cip_dict.items()]

    @property
    def get_evaluation_list(self) -> list[EvaluationInfo]:
        evaluation_title_list = self._get_value_list("Titre de l'évaluation", col_offset=3, vertical_search=False)
        maximum_grade_list = self._get_value_list("Note sur:", col_offset=3, vertical_search=False)

        evaluation_info_list = []

        for index, evaluation in enumerate(evaluation_title_list, start=1):
            evaluation_numbered_title = f"{index} - {evaluation.strip()}"
            maximum_grade = decimal_helpers.safe_decimal(maximum_grade_list[index - 1])
            evaluation_info_list.append(EvaluationInfo(title=evaluation_numbered_title,
                                                       maximum_grade=maximum_grade))
        return evaluation_info_list

    def _find_value_by_text_coord(self, row_text_value: str, col_text_value: str) -> Any | None:
        cell_row = None
        cell_col = None

        for row in self.ws.iter_rows():
            for cell in row:
                if cell.value == row_text_value:
                    cell_row = cell.row

                if cell.value == col_text_value:
                    cell_col = cell.column

                if cell_row and cell_col:
                    return self.ws.cell(row=cell_row, column=cell_col).value

        return None

    def _find_value_by_text_offset(self, text: str, row_offset: int = 0, col_offset: int = 0) -> str | None:
        cell_row = None
        cell_col = None

        for row in self.ws.iter_rows():
            for cell in row:
                if cell.value == text:
                    cell_row = cell.row
                    cell_col = cell.column

                if cell_row and cell_col:
                    return self.ws.cell(row=cell_row + row_offset, column=cell_col + col_offset).value  # type: ignore

        return None

    def _get_value_list(self, text: str, row_offset: int = 0, col_offset: int = 0,
                        vertical_search: bool = True) -> list:
        values = []

        for row in self.ws.iter_rows():
            for cell in row:
                if cell.value == text:
                    starting_row = cell.row
                    starting_col = cell.column

                    while True:
                        current_cell = self.ws.cell(row=starting_row + row_offset, column=starting_col + col_offset)
                        value = current_cell.value

                        if value is None or str(value).strip() == "":
                            break

                        values.append(value)

                        if vertical_search:
                            starting_row += 1
                        else:
                            starting_col += 1

        return values