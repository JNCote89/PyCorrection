import logging
from pathlib import Path
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from src.py_correction.core.domains.grade.grade_dtos import GeNoteTransactionDTO
from src.py_correction.engine.excel import excel_operations
from src.py_correction.engine.helpers import decimal_helpers

logger = logging.getLogger(__name__)


class GeNoteGradeFileOperations:

    def __init__(self, file_path: str | Path):
        self.file_path = file_path

    def add_grades(self, genote_transaction_dtos: list[GeNoteTransactionDTO]) -> list[str] | None:
        failed_operations = []

        with excel_operations.safe_load_workbook(excel_path=self.file_path, data_only=True) as wb:
            if wb is None:
                logger.warning(f"""La feuille GeNote dans le répertoire {self.file_path} n'a pas pu être modifiée.
                                   Impossible d'ajouter des notes. Veuillez vérifier l'intégrité du fichier.""")
                return None

            ws = wb[wb.sheetnames[0]]

            for transaction_dto in genote_transaction_dtos:

                row_idx, col_idx = self._find_coordinates(ws=ws, student_name=transaction_dto.student_name,
                                                          evaluation_title=transaction_dto.evaluation_title)

                if row_idx and col_idx:
                    ws.cell(row=row_idx, column=col_idx).value = decimal_helpers.safe_decimal(
                        transaction_dto.student_grade)
                    logger.info(f"La note de {transaction_dto.student_grade} a été ajoutée avec succès pour l'étudiant "
                                f"{transaction_dto.student_name} pour l'évaluation {transaction_dto.evaluation_title}")
                else:
                    failed_operations.append(transaction_dto.student_name)
                    logger.info(f"La note pour {transaction_dto.student_name} n'a pas pu être ajouté à la feuille "
                                   f"GeNote pour l'évaluation {transaction_dto.evaluation_title}")

            wb.save(self.file_path)

            return failed_operations

    @staticmethod
    def _find_coordinates(ws: Worksheet, student_name: str, evaluation_title: str) -> tuple[Any, Any]:
        student_row = None
        evaluation_col = None

        for row in ws.iter_rows():
            for cell in row:
                val = cell.value

                if not student_row and val == student_name:
                    student_row = cell.row

                if not evaluation_col and isinstance(val, str):
                    if val == evaluation_title or val in evaluation_title: # type: ignore[arg-type]
                        evaluation_col = cell.column

                if student_row and evaluation_col:
                    return student_row, evaluation_col

        return student_row, evaluation_col