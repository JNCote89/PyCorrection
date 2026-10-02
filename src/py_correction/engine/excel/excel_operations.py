from collections import Counter
from contextlib import closing, contextmanager
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

logger = logging.getLogger(__name__)

MAX_KEYWORD_COUNT = 6


@contextmanager
def safe_load_workbook(excel_path: Path | str, **kwargs):
    if isinstance(excel_path, str):
        excel_path = Path(excel_path).resolve()

    if not isinstance(excel_path, Path) or not excel_path.is_file():
        logger.warning(f"Excel file path does not exist or is invalid: {excel_path}")
        yield None
        return

    try:
        wb = load_workbook(excel_path, **kwargs)

    except (FileNotFoundError, PermissionError, InvalidFileException, ValueError, KeyError) as e:
        logger.warning(f"Failed to load or read Excel file at {excel_path}: {e}")
        yield None
        return

    with closing(wb):
        yield wb


def get_excel_paths_from_directory(directory: Path | None) -> list[Path] | None:
    if directory is None:
        return None
    return [file for file in directory.rglob("*") if file.suffix.lower() in (".xlsx", ".xls")]


def get_excel_sheet_names(excel_path: Path | None) -> list[str] | None:
    if excel_path is None:
        return None
    with safe_load_workbook(excel_path, read_only=True, data_only=True) as workbook:
        if workbook is None:
            return None
        return workbook.sheetnames


def get_excel_sheet_keyword_options(excel_path: Path | None, sheet_name: str | None) -> list[str] | None:
    if excel_path:
        with safe_load_workbook(excel_path, read_only=True, data_only=True) as workbook:
            if workbook and sheet_name:
                ws = workbook[sheet_name]

                word_counter = Counter()

                for row in ws.iter_rows(values_only=True):
                    for val in row:
                        if isinstance(val, str):
                            word_count = len(val.split())
                            if 1 <= word_count <= MAX_KEYWORD_COUNT:
                                cleaned_val = val.strip()
                                word_counter[cleaned_val] += 1

                unique_keywords = [word for word, count in word_counter.items() if count == 1]

                return sorted(unique_keywords)
            else:
                return None
    else:
        return None


def exclude_duplicate_keyword(row_keyword: str | None, keywords: list[str] | None) -> list[str] | None:
    if isinstance(keywords, list):
        return [keyword for keyword in keywords if keyword != row_keyword]

    return None
